#!/usr/bin/env python3
"""Builds the final Oxyderm Client Marketing Sheet from merged_clients.json."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/Users/genesis/Desktop/Oxyderm-Build/data/form_extraction"
OUT = f"{BASE}/Oxyderm_Client_Marketing_Sheet.xlsx"

with open(f"{BASE}/merged_clients.json") as f:
    rows = json.load(f)

TOTAL_UNIQUE_CLIENTS = 1652
TOTAL_EXCLUDED_NO_CONTACT = 551

# Sort by visit count descending — highest-value / most engaged clients first
rows.sort(key=lambda r: r["treatment_visit_count"], reverse=True)

wb = Workbook()
ws = wb.active
ws.title = "Client Marketing List"

headers = [
    "Name", "Phone", "Email", "Total Visits", "Treatment Areas",
    "Membership Notes", "Ethnicity", "Skin Type #", "First Form Date",
    "Last Form Date", "Total Form Submissions", "Square Customer ID"
]
ws.append(headers)

header_fill = PatternFill("solid", fgColor="00172D")  # Oxyderm navy
header_font = Font(bold=True, color="FFFFFF", size=11)
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(*(Side(style="thin", color="DDDDDD"),) * 4)

for r in rows:
    ws.append([
        r["name"], r["phone"], r["email"], r["treatment_visit_count"],
        r["treatments"], r["membership_notes"], r["ethnicity"],
        r["skin_type_number"], r["first_form_date"], r["last_form_date"],
        r["total_form_submissions"], r["square_customer_id"],
    ])

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=(cell.column_letter in ("E", "F")))

# Column widths tuned for readability
widths = {"A": 22, "B": 16, "C": 28, "D": 12, "E": 55, "F": 30, "G": 14, "H": 12, "I": 14, "J": 14, "K": 10, "L": 28}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

# Highlight high-value clients (10+ visits) for quick marketing targeting
high_value_fill = PatternFill("solid", fgColor="FFF3D6")  # soft gold tint
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    visits_cell = row[3]  # column D
    if isinstance(visits_cell.value, int) and visits_cell.value >= 10:
        for cell in row:
            cell.fill = high_value_fill

# Summary sheet
summary = wb.create_sheet("Summary")
summary_rows = [
    ["Oxyderm Client Marketing List — Summary", ""],
    ["", ""],
    ["Generated from", "8,011 historical form emails (Treatment Records, Skin Chart, Medical History)"],
    ["Total unique clients found (deduped by name)", TOTAL_UNIQUE_CLIENTS],
    ["Clients with verified contact info (in this sheet)", len(rows)],
    ["Excluded (no phone/email match in Square)", TOTAL_EXCLUDED_NO_CONTACT],
    ["High-value clients (10+ visits, highlighted gold)", sum(1 for r in rows if r["treatment_visit_count"] >= 10)],
    ["", ""],
    ["IMPORTANT: Medical History Form data was NOT included.", "Only name + date were extracted from that source; all health/medical fields were deliberately excluded for privacy/compliance."],
    ["Contact info source", "Matched against live Square customer records by name. Clients with no Square match were excluded entirely, per instruction."],
]
for row in summary_rows:
    summary.append(row)
summary.column_dimensions["A"].width = 45
summary.column_dimensions["B"].width = 70
summary["A1"].font = Font(bold=True, size=14, color="00172D")
for r in range(3, len(summary_rows) + 1):
    summary.cell(row=r, column=1).font = Font(bold=True)
    summary.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print(f"Saved: {OUT}")
print(f"Rows: {len(rows)}")
